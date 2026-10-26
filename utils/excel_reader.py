from pathlib import Path
from openpyxl import load_workbook

DATA_FILE = Path(__file__).resolve().parents[1] / "data" / "test_data.xlsx"

def read_test_cases():
    """Read enabled rows from the TestCases worksheet as dictionaries."""
    wb = load_workbook(DATA_FILE, data_only=True)
    ws = wb["TestCases"]
    headers = [cell.value for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if str(row.get("enabled", "Y")).strip().upper() == "Y":
            rows.append(row)
    wb.close()
    return rows
